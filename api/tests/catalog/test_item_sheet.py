"""Items to and from a spreadsheet: export, edit, import back by code."""

from __future__ import annotations

import io
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.db import SessionLocal
from app.models import Item
from tests.catalog.test_api_import import CatalogEnv

X = "/api/item-sheet"


@pytest.fixture
def shop(catalog_client: CatalogEnv) -> tuple[TestClient, dict[str, str], list[dict]]:
    client, h, _ = catalog_client
    cat = client.post("/api/item-categories", headers=h, json={"name": "Drinkware"}).json()
    items = []
    for name, sku, rate in (("Beer Mug", "100500", "1256"), ("Juice Glass", "100501", "190")):
        r = client.post(
            "/api/items",
            headers=h,
            json={
                "name": name,
                "uom": "nos",
                "default_rate": rate,
                "sku": sku,
                "category_id": cat["id"],
            },
        )
        assert r.status_code == 201, r.text
        items.append(r.json())
    return client, h, items


def _sheet(rows: list[list[Any]], header: list[str] | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(
        header or ["Code", "Name", "Group", "Availability", "Selling price", "Pack qty", "HSN"]
    )
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _post(client: TestClient, h: dict[str, str], data: bytes, name="items.xlsx", **q: Any) -> Any:
    return client.post(
        f"{X}/import", headers=h, params=q, files={"file": (name, data, "application/octet-stream")}
    )


def _get(item_id: str) -> Item:
    with SessionLocal() as s:
        it = s.scalar(select(Item).where(Item.id == item_id))
        s.expunge(it)
        return it


def test_export_has_the_items_and_their_codes(shop) -> None:
    client, h, items = shop
    r = client.post(f"{X}/export", headers=h, json={})
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    ws = load_workbook(io.BytesIO(r.content)).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][:5] == ("Code", "Name", "Group", "Availability", "Selling price")
    by_code = {r[0]: r for r in rows[1:]}
    assert by_code["100500"][1] == "Beer Mug" and by_code["100500"][2] == "Drinkware"
    assert by_code["100500"][3] == "In stock" and by_code["100500"][4] == 1256
    only = client.post(f"{X}/export", headers=h, json={"ids": [items[1]["id"]]})
    assert len(list(load_workbook(io.BytesIO(only.content)).active.iter_rows())) == 2
    flt = client.post(f"{X}/export", headers=h, json={"filter": {"q": "juice"}})
    assert len(list(load_workbook(io.BytesIO(flt.content)).active.iter_rows())) == 2


def test_import_dry_run_then_apply(shop) -> None:
    client, h, items = shop
    data = _sheet(
        [
            ["100500", "Beer Mug Large", "", "Out of stock", 1300, 6, "7013"],
            ["100501", "", "", "", "", "", ""],  # blank cells change nothing
        ]
    )
    pre = _post(client, h, data).json()  # dry run is the default
    assert pre["dry_run"] and pre["changed"] == 1 and pre["unchanged"] == 1 and pre["errors"] == 0
    assert _get(items[0]["id"]).name == "Beer Mug"
    done = _post(client, h, data, dry_run="false").json()
    assert done["changed"] == 1
    it = _get(items[0]["id"])
    assert (it.name, it.availability, float(it.default_rate), it.pack_qty, it.hsn_code) == (
        "Beer Mug Large",
        "out_of_stock",
        1300.0,
        6,
        "7013",
    )
    assert _get(items[1]["id"]).name == "Juice Glass"


def test_import_reports_each_bad_row_and_applies_the_rest(shop) -> None:
    client, h, items = shop
    data = _sheet(
        [
            ["999999", "Ghost", "", "", "", "", ""],
            ["100500", "Juice Glass", "", "", "", "", ""],  # name clashes with the other item
            ["100501", "Juice Tumbler", "Nowhere", "", "", "", ""],  # unknown group
            ["100501", "", "", "Sold out maybe", "", "", ""],  # repeated code
            ["100500", "", "", "Expected", -5, "", ""],  # repeated code, never reached
        ]
    )
    r = _post(client, h, data, dry_run="false").json()
    details = {row["row"]: row["detail"] for row in r["rows"]}
    assert r["errors"] == 5 and r["changed"] == 0
    assert "no item has this code" in details[2]
    assert "existing item" in details[3]
    assert "unknown group" in details[4]
    assert "twice" in details[5] and "twice" in details[6]
    # one good row alongside bad ones still goes through
    ok = _sheet(
        [
            ["999999", "x", "", "", "", "", ""],
            ["100501", "Juice Tumbler", "Drinkware", "", 200, "", ""],
        ]
    )
    r = _post(client, h, ok, dry_run="false").json()
    assert r["changed"] == 1 and r["errors"] == 1
    assert _get(items[1]["id"]).name == "Juice Tumbler"


def test_import_validates_values(shop) -> None:
    client, h, _ = shop
    data = _sheet(
        [
            ["100500", "", "", "Maybe", "", "", ""],
            ["100501", "", "", "", "abc", "", ""],
        ]
    )
    r = _post(client, h, data).json()
    assert r["errors"] == 2
    bad = _sheet([["100500", "", "", "", "", 2.5, ""]])
    assert "whole number" in _post(client, h, bad).json()["rows"][0]["detail"]
    bad = _sheet([["100500", "", "", "", "", "", "12"]])
    assert "HSN" in _post(client, h, bad).json()["rows"][0]["detail"]


def test_csv_works_and_bad_files_are_refused(shop) -> None:
    client, h, items = shop
    csv_bytes = "\ufeffCode,Name,Selling price\n100500,Beer Mug,1500\n".encode()
    r = _post(client, h, csv_bytes, name="items.csv", dry_run="false")
    assert r.status_code == 200 and r.json()["changed"] == 1
    assert float(_get(items[0]["id"]).default_rate) == 1500.0
    assert _post(client, h, b"not a workbook").status_code == 422
    assert _post(client, h, _sheet([], header=["Name", "Price"])).status_code == 422
    assert _post(client, h, b"").status_code == 422


def test_search_finds_an_item_by_its_code(shop) -> None:
    client, h, items = shop
    r = client.get("/api/items", headers=h, params={"q": "100500"})
    assert [i["id"] for i in r.json()] == [items[0]["id"]]
    r = client.get("/api/items", headers=h, params={"q": "10050"})
    assert {i["id"] for i in r.json()} == {items[0]["id"], items[1]["id"]}


def test_active_jobs_lists_running_work(shop) -> None:
    client, h, items = shop
    assert client.get("/api/jobs/active", headers=h).json() == []
    from app.models import CatalogOutputJob, Tenant

    with SessionLocal() as s:
        tid = s.scalar(select(Tenant.id))
        s.add(
            CatalogOutputJob(
                tenant_id=tid, kind="item_labels", status="running", progress=40, total=120
            )
        )
        s.add(
            CatalogOutputJob(tenant_id=tid, kind="item_labels", status="done", progress=1, total=1)
        )
        s.commit()
    jobs = client.get("/api/jobs/active", headers=h).json()
    assert len(jobs) == 1
    assert (
        jobs[0]["title"] == "Printing labels"
        and jobs[0]["progress"] == 40
        and jobs[0]["total"] == 120
    )
